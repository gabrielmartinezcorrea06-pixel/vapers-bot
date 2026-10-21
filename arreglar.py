from openpyxl import load_workbook 
wb = load_workbook('ventas_vapers.xlsx') 
ws = wb['Ventas'] 
datos = [] 
for row in ws.iter_rows(min_row=2, values_only=True): 
    if any(v is not None for v in row): 
        datos.append(row) 
del wb['Ventas'] 
ws2 = wb.create_sheet('Ventas', 0) 
for i, fila in enumerate(datos, 2): 
    for j, val in enumerate(fila, 1): 
        ws2.cell(row=i, column=j, value=val) 
wb.save('ventas_vapers.xlsx') 
print("Listo:", len(datos)) 
